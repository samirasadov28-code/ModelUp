"""
Inject questionnaire answers into the Scen sheet of the source Excel models.

Cell reference map (based on StartUp Model.xlsx Scen sheet):
  Scen!B3   — Model start date
  Scen!B6   — Development stage length (months)
  Scen!B7   — Operation stage length (years)
  Scen!B10  — Initial users
  Scen!B11  — Max capacity
  Scen!B12  — User growth curve (1=Base, 2=Aggressive, 3=Conservative)
  Scen!B16  — Tier 1 name
  Scen!C16  — Tier 1 % allocation
  Scen!D16  — Tier 1 monthly price
  Scen!B17  — Tier 2 name
  Scen!C17  — Tier 2 % allocation
  Scen!D17  — Tier 2 monthly price
  Scen!B18  — Tier 3 name / Scen!B19 — Tier 4 name
  Scen!B22  — Tier 1 monthly churn %
  Scen!B23  — Tier 2 monthly churn %
  Scen!B24  — Tier 3 monthly churn %
  Scen!B25  — Tier 4 monthly churn %
  Scen!B30  — Equity injection amount
  Scen!B31  — Equity injection timing (month)
  Scen!B35  — Primary currency (USD/GBP/EUR)
  Scen!B38  — CPI/inflation %
"""

import os
import shutil
import tempfile
from datetime import date
from typing import Any, Dict, Optional

import openpyxl


MODELS_DIR = os.path.join(os.path.dirname(__file__), "..", "source_models")

MODEL_FILES = {
    "saas": "StartUp Model.xlsx",
    "alternative": "StartUp Model Alternative.xlsx",
    "project_finance": "StartUp Model.xlsx",  # Falls back; PF model exported as xlsx
}

CURRENCY_MAP = {
    "us": "USD",
    "uk": "GBP",
    "eu": "EUR",
    "asia": "USD",
    "global": "USD",
}

CURVE_MAP = {
    "base": 1,
    "aggressive": 2,
    "conservative": 3,
}


def get_source_path(model_type: str) -> str:
    filename = MODEL_FILES.get(model_type, "StartUp Model.xlsx")
    return os.path.join(MODELS_DIR, filename)


def inject_answers(answers: Dict[str, Any], model_type: str) -> str:
    """
    Copy the source model to a temp file, inject answers, return the temp path.
    Caller is responsible for deleting the temp file.
    """
    source = get_source_path(model_type)

    if not os.path.exists(source):
        raise FileNotFoundError(
            f"Source model not found: {source}. "
            "Place the Excel source files in excel-service/source_models/"
        )

    tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    tmp.close()
    shutil.copy2(source, tmp.name)

    wb = openpyxl.load_workbook(tmp.name, keep_vba=False)

    # Prefer sheet named "Scen"; fall back to first sheet
    sheet_name = "Scen" if "Scen" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]

    # ── Timing ────────────────────────────────────────────────────────────────
    start_date = answers.get("modelStartDate") or date.today().isoformat()
    _safe_write(ws, "B3", start_date)
    _safe_write(ws, "B6", 3)   # default 3-month dev stage
    _safe_write(ws, "B7", 3)   # 3-year operation stage

    # ── Users ─────────────────────────────────────────────────────────────────
    _safe_write(ws, "B10", answers.get("year1UserTarget", 100))
    _safe_write(ws, "B12", CURVE_MAP.get(answers.get("growthCurve", "base"), 1))

    # ── Tiers ─────────────────────────────────────────────────────────────────
    tiers = answers.get("tiers", [])
    tier_rows = [16, 17, 18, 19]
    churn_rows = [22, 23, 24, 25]
    monthly_churn = answers.get("monthlyChurnRate", 5) / 100

    for i, row in enumerate(tier_rows):
        if i < len(tiers):
            t = tiers[i]
            _safe_write(ws, f"B{row}", t.get("name", f"Tier {i+1}"))
            alloc = t.get("allocationPercent", 0) / 100
            _safe_write(ws, f"C{row}", alloc)
            _safe_write(ws, f"D{row}", t.get("monthlyPrice", 0))
        _safe_write(ws, f"B{churn_rows[i]}", monthly_churn)

    # ── Fundraising ───────────────────────────────────────────────────────────
    _safe_write(ws, "B30", answers.get("fundingAsk", 0))
    _safe_write(ws, "B31", 1)  # inject at month 1

    # ── Currency / Macro ──────────────────────────────────────────────────────
    currency = CURRENCY_MAP.get(answers.get("geography", "us"), "USD")
    _safe_write(ws, "B35", currency)
    _safe_write(ws, "B38", 0.02)  # 2% CPI default

    wb.save(tmp.name)
    wb.close()
    return tmp.name


def _safe_write(ws: Any, cell: str, value: Any) -> None:
    try:
        ws[cell] = value
    except Exception:
        pass
